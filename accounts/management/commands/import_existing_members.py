from collections import defaultdict
from math import isfinite
from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from accounts.models import ExistingMember, UserProfile
from accounts.utils import normalize_nigerian_phone_number


class Command(BaseCommand):
    help = "Import existing-member phone numbers into the account activation roster."

    def add_arguments(self, parser):
        parser.add_argument(
            "--file",
            default=str(Path(settings.BASE_DIR) / "Group Numbers.xlsx"),
            help="Excel workbook path (defaults to Group Numbers.xlsx in the project root).",
        )
        parser.add_argument(
            "--sheet",
            default="All Contacts",
            help="Worksheet name (defaults to All Contacts).",
        )
        parser.add_argument(
            "--phone-column",
            help="Phone column header or one-based column number; detected from the header by default.",
        )
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Validate and report counts without importing member records.",
        )
        parser.add_argument(
            "--skip-invalid",
            action="store_true",
            help="Skip invalid phone-number rows and report their row numbers.",
        )

    def handle(self, *args, **options):
        file_path = Path(options["file"])
        try:
            workbook = load_workbook(file_path, read_only=True, data_only=True)
        except (OSError, InvalidFileException) as exc:
            raise CommandError(f"Could not read Excel workbook: {file_path}") from exc

        try:
            if options["sheet"] not in workbook.sheetnames:
                raise CommandError(f"Worksheet not found: {options['sheet']}")

            worksheet = workbook[options["sheet"]]
            rows = worksheet.iter_rows(values_only=True)
            header = next(rows, None)
            if header is None:
                raise CommandError("The selected worksheet is empty.")

            phone_column = self._phone_column(header, options["phone_column"])
            phone_rows = defaultdict(list)
            invalid_rows = []
            for row_number, row in enumerate(rows, start=2):
                value = row[phone_column] if phone_column < len(row) else None
                if value is None or (isinstance(value, str) and not value.strip()):
                    continue
                try:
                    normalized = normalize_nigerian_phone_number(
                        self._cell_text(value)
                    )
                except (ValueError, TypeError):
                    invalid_rows.append(row_number)
                    continue
                phone_rows[normalized].append(row_number)

            duplicates = [
                row_numbers
                for row_numbers in phone_rows.values()
                if len(row_numbers) > 1
            ]
            if duplicates:
                duplicate_rows = sorted(
                    row_number for row_numbers in duplicates for row_number in row_numbers
                )
                raise CommandError(
                    "Duplicate phone numbers after normalization in spreadsheet rows: "
                    + ", ".join(map(str, duplicate_rows))
                )
            if invalid_rows and not options["skip_invalid"]:
                raise CommandError(
                    "Invalid Nigerian phone numbers in spreadsheet rows: "
                    + ", ".join(map(str, invalid_rows[:50]))
                )

            phone_numbers = list(phone_rows)
            if not phone_numbers:
                raise CommandError("No Nigerian phone numbers found in the worksheet.")

            if options["dry_run"]:
                self.stdout.write(
                    f"Validated {len(phone_numbers)} unique phone numbers; "
                    f"{len(invalid_rows)} invalid rows skipped."
                )
                if invalid_rows:
                    self.stdout.write(
                        "Skipped worksheet rows: "
                        + ", ".join(map(str, invalid_rows[:50]))
                    )
                return

            existing = set(
                str(phone)
                for phone in ExistingMember.objects.filter(
                    phone_number__in=phone_numbers
                ).values_list("phone_number", flat=True)
            )
            profile_matches = set(
                str(phone)
                for phone in UserProfile.objects.filter(
                    phone_number__in=phone_numbers
                ).values_list("phone_number", flat=True)
            )
            to_import = [
                phone for phone in phone_numbers if phone not in existing
            ]

            with transaction.atomic():
                imported = sum(
                    ExistingMember.objects.get_or_create(phone_number=phone)[1]
                    for phone in to_import
                )

            self.stdout.write(
                self.style.SUCCESS(
                    f"Imported {imported} member records; "
                    f"{len(existing)} were already present; "
                    f"{len(profile_matches)} match existing profiles; "
                    f"{len(invalid_rows)} invalid rows skipped."
                )
            )
            if invalid_rows:
                self.stdout.write(
                    "Skipped worksheet rows: " + ", ".join(map(str, invalid_rows[:50]))
                )
        finally:
            workbook.close()

    @staticmethod
    def _phone_column(header, requested_column):
        if requested_column:
            if requested_column.isdigit():
                column = int(requested_column) - 1
                if 0 <= column < len(header):
                    return column
                raise CommandError("Phone column number is outside the worksheet.")
            for index, value in enumerate(header):
                if str(value or "").strip().casefold() == requested_column.strip().casefold():
                    return index
            raise CommandError(f"Phone column header not found: {requested_column}")

        for index, value in enumerate(header):
            normalized_header = str(value or "").strip().casefold()
            if "phone" in normalized_header or "mobile" in normalized_header:
                return index
        raise CommandError(
            "Could not detect a phone column; specify --phone-column."
        )

    @staticmethod
    def _cell_text(value):
        if isinstance(value, bool):
            raise ValueError("Boolean values are not phone numbers.")
        if isinstance(value, int):
            return str(value)
        if isinstance(value, float):
            if not isfinite(value) or not value.is_integer():
                raise ValueError("Phone number cell must contain an integer.")
            return str(int(value))
        return str(value)
